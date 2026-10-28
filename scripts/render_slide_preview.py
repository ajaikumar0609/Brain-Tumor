#!/usr/bin/env python3
"""Shape-based APPROXIMATE preview of IA3_Presentation.pptx (PIL only; no PowerPoint/Keynote/LibreOffice).

Draws rectangles, text boxes (word-wrapped, per-run size/bold/colour), tables and pictures from the stored shape geometry.
Calibri is not installed, so Arial (about 10% wider) is used: the preview is deliberately slightly pessimistic about wrapping.
It is NOT a PowerPoint/Keynote render and does not replace opening the deck in one of them.

    python scripts/render_slide_preview.py
"""
import base64, io, json, sys
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
PPTX = ROOT / "outputs/submission/IA3_Presentation.pptx"
OUT = ROOT / "outputs/submission/slide_preview"
PX = 144  # pixels per inch -> 1920 x 1080
IN = 914400
FONT = {False: "/System/Library/Fonts/Supplemental/Arial.ttf", True: "/System/Library/Fonts/Supplemental/Arial Bold.ttf"}
_cache = {}


def font(pt, bold):
    k = (round(pt * PX / 72), bold)
    if k not in _cache:
        _cache[k] = ImageFont.truetype(FONT[bool(bold)], k[0])
    return _cache[k]


def px(emu):
    return emu / IN * PX


def rgb(c, default):
    try:
        return tuple(c.rgb) if c is not None and c.type is not None else default
    except Exception:
        return default


def wrap_runs(paras, width):
    """paras: list of (runs[(text,pt,bold,color)], space_after_pt). Return laid-out lines [(y_offset, [(x,text,font,color)], line_h)] and total height."""
    lines, y = [], 0.0
    for runs, space_after in paras:
        words = []
        for text, pt, bold, col in runs:
            for i, w in enumerate(text.split(" ")):
                words.append((w, pt, bold, col))
        cur, curw, maxpt = [], 0.0, max([r[1] for r in runs] + [10])
        def flush():
            nonlocal cur, curw, y
            lh = maxpt * 1.2 * PX / 72
            lines.append((y, cur, lh)); y += lh; cur, curw = [], 0.0
        for w, pt, bold, col in words:
            f = font(pt, bold); ww = f.getlength(w + " ")
            if cur and curw + f.getlength(w) > width:
                flush()
            cur.append((curw, w + " ", f, col)); curw += ww
        flush()
        y += space_after * PX / 72
    return lines, y


def draw_text(d, x, y, w, h, tf, anchor_mid, flags, label, inset=(0.1, 0.05), align_default=None):
    paras = []
    for p in tf.paragraphs:
        runs = []
        for r in p.runs:
            f = r.font
            runs.append((r.text, f.size.pt if f.size else 18, bool(f.bold), rgb(f.color, (0x33, 0x33, 0x33)) if f.color and f.color.type else (0x33, 0x33, 0x33)))
        if not runs:
            runs = [("", 12, False, (0, 0, 0))]
        sa = p.space_after.pt if p.space_after is not None else 0
        paras.append((runs, sa))
    ix, iy = inset[0] * PX, inset[1] * PX
    lines, total = wrap_runs(paras, w - 2 * ix)
    oy = y + iy
    if anchor_mid:
        oy = y + (h - total) / 2
    for ly, items, lh in lines:
        for xo, text, f, col in items:
            d.text((x + ix + xo, oy + ly), text, font=f, fill=col)
    if total + 2 * iy > h + 0.05 * PX:
        flags.append(f"text taller than its box by {round((total + 2 * iy - h) / PX, 2)} in ({label})")
    # text running past slide edges
    if oy + total > 7.5 * PX + 1 or x + w > 13.34 * PX:
        flags.append(f"content beyond the slide edge ({label})")
    return total


def render_slide(prs, slide, n):
    W, H = int(prs.slide_width / IN * PX), int(prs.slide_height / IN * PX)
    im = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(im)
    flags, inventory = [], []
    for sh in slide.shapes:
        x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
        kind = str(sh.shape_type).split(".")[-1].split(" ")[0]
        inventory.append(kind)
        if sh.shape_type == 13:
            pic = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
            im.paste(pic.resize((max(1, int(w)), max(1, int(h))), Image.LANCZOS), (int(x), int(y)))
            ar_img, ar_box = pic.size[0] / pic.size[1], w / h
            if abs(ar_img - ar_box) / ar_box > 0.03:
                flags.append(f"picture aspect ratio distorted ({ar_img:.2f} vs {ar_box:.2f})")
            continue
        if getattr(sh, "has_table", False) and sh.has_table:
            tb, cy = sh.table, y
            for r in tb.rows:
                cx, rowh, cells = x, r.height / IN * PX, []
                for j, c in enumerate(r.cells):
                    cw = tb.columns[j].width / IN * PX
                    runs = [(rr.text, rr.font.size.pt if rr.font.size else 12, bool(rr.font.bold), rgb(rr.font.color, (0x33, 0x33, 0x33))) for p in c.text_frame.paragraphs for rr in p.runs] or [("", 12, False, (0, 0, 0))]
                    lines, th = wrap_runs([(runs, 0)], cw - 0.12 * PX)
                    rowh = max(rowh, th + 0.06 * PX)
                    cells.append((cx, cw, c, lines))
                    cx += cw
                for cx, cw, c, lines in cells:
                    fillc = rgb(c.fill.fore_color, (255, 255, 255)) if c.fill.type == 1 else (255, 255, 255)
                    d.rectangle([cx, cy, cx + cw, cy + rowh], fill=fillc, outline=(200, 200, 200))
                    for ly, items, lh in lines:
                        for xo, text, f, col in items:
                            d.text((cx + 0.06 * PX + xo, cy + 0.03 * PX + ly), text, font=f, fill=col)
                cy += rowh
            if cy > y + h + 0.05 * PX:
                flags.append(f"table grows from {round(h / PX, 2)} to {round((cy - y) / PX, 2)} in (tables grow downward)")
            if cy > H - 0.4 * PX and n != 1:
                flags.append("table reaches the footer area")
            continue
        fillc = outl = None
        try:
            if sh.fill.type == 1:
                fillc = tuple(sh.fill.fore_color.rgb)
        except Exception:
            pass
        try:
            if sh.line.fill.type == 1:
                outl = tuple(sh.line.color.rgb)
        except Exception:
            pass
        if fillc or outl:
            d.rectangle([x, y, x + w, y + h], fill=fillc, outline=outl, width=2 if outl else 0)
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            mid = str(tf.vertical_anchor).startswith("MIDDLE")
            ins = (tf.margin_left / IN if tf.margin_left is not None else 0.1, tf.margin_top / IN if tf.margin_top is not None else 0.05)
            draw_text(d, x, y, w, h, tf, mid, flags, tf.text[:30].replace("\n", " "), ins)
    return im, flags, inventory


def main():
    prs = Presentation(PPTX)
    OUT.mkdir(parents=True, exist_ok=True)
    report, imgs = [], []
    for n, s in enumerate(prs.slides, 1):
        im, flags, inv = render_slide(prs, s, n)
        im.save(OUT / f"slide_{n:02d}.png")
        imgs.append(im)
        title = next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() and sh.top < 0.5 * IN * 2 and n > 1), "Title slide")
        report.append({"slide": n, "title": title[:90], "elements": {k: inv.count(k) for k in set(inv)}, "flags": flags})
    # contact sheet 3 x 5
    tw, th, gap = 640, 360, 14
    cols, rows = 3, 5
    sheet = Image.new("RGB", (cols * tw + (cols + 1) * gap, rows * (th + 28) + gap), (235, 235, 235))
    dd = ImageDraw.Draw(sheet)
    for i, im in enumerate(imgs):
        cx, cy = gap + (i % cols) * (tw + gap), gap + (i // cols) * (th + 28)
        sheet.paste(im.resize((tw, th), Image.LANCZOS), (cx, cy + 22))
        dd.text((cx, cy + 2), f"Slide {i + 1}", font=font(11, True), fill=(40, 40, 40))
    sheet.save(OUT / "contact_sheet.png")
    json.dump(report, open(OUT / "approx_flags.json", "w"), indent=2)
    # index.html (self-contained: images embedded as base64)
    def b64(p):
        return base64.b64encode((OUT / p).read_bytes()).decode()
    slides_html = "\n".join(f'<section id="s{r["slide"]}"><h2>Slide {r["slide"]}: {r["title"]}</h2><img src="data:image/png;base64,{b64(f"slide_{r["slide"]:02d}.png")}" alt="slide {r["slide"]}"><p class="fl">{"; ".join(r["flags"]) or "no automatic flags"}</p></section>' for r in report)
    nav = " ".join(f'<a href="#s{r["slide"]}">{r["slide"]}</a>' for r in report)
    html = f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>IA3 slide preview (approximation)</title>
<style>body{{font-family:-apple-system,Helvetica,Arial,sans-serif;margin:0;background:#f2f2f2;color:#222}}header{{position:sticky;top:0;background:#1f3a5f;color:#fff;padding:8px 16px;z-index:2}}header a{{color:#f2c97a;margin-right:8px;text-decoration:none}}
.warn{{background:#fdf1dc;border:1px solid #c98a1a;padding:10px 16px;margin:16px}}section{{margin:16px;background:#fff;padding:12px;border-radius:6px}}img{{max-width:100%;border:1px solid #ccc}}.fl{{font-size:13px;color:#555}}</style></head><body>
<header><strong>IA3 slide preview</strong> &nbsp; <a href="#contact">contact sheet</a> {nav}</header>
<div class="warn"><strong>Approximation only.</strong> Drawn from the PPTX shape geometry with PIL and Arial (Calibri is not installed). This is not a PowerPoint or Keynote render and does not replace opening the deck in one of them.</div>
<section id="contact"><h2>Contact sheet</h2><img src="data:image/png;base64,{b64("contact_sheet.png")}" alt="contact sheet"></section>
{slides_html}</body></html>"""
    (OUT / "index.html").write_text(html)
    for r in report:
        print(r["slide"], r["title"][:50], "|", "; ".join(r["flags"]) or "-")


if __name__ == "__main__":
    main()
